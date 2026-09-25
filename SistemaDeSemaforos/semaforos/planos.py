"""Leitura dos tempos semafóricos da planilha recebida."""

import openpyxl


def read_plans(path):
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    crossings = {}
    for row in range(1, sheet.max_row + 1):
        name = sheet.cell(row, 1).value
        if not isinstance(name, str) or " x " not in name:
            continue
        stages = {}
        for stage in "ABCD":
            description = sheet.cell(row + 12 + "ABCD".index(stage), 2).value
            if description:
                stages[stage] = str(description).strip()
        plans = {}
        for plan_row in range(row + 4, row + 8):
            plan_id = sheet.cell(plan_row, 1).value
            if not isinstance(plan_id, (int, float)):
                continue
            phases = {}
            for column, stage in ((3, "A"), (7, "B"), (11, "C"), (15, "D")):
                values = [sheet.cell(plan_row, column + offset).value for offset in range(3)]
                if all(isinstance(value, (int, float)) for value in values):
                    phases[stage] = dict(zip(("green", "yellow", "clearance"), values))
            cycle = sheet.cell(plan_row, 19).value
            calculated = sum(sum(phase.values()) for phase in phases.values())
            if abs(calculated - cycle) > 0.01:
                raise ValueError(f"Ciclo inconsistente em {name}, plano {plan_id}")
            plans[str(int(plan_id))] = {
                "offset": sheet.cell(plan_row, 2).value,
                "cycle": cycle,
                "stages": phases,
            }
        crossings[name] = {"stages": stages, "plans": plans, "sheet_row": row}
    return crossings
